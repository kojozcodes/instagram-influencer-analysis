import io

from openpyxl import load_workbook

from app.excel import HEADERS, build_excel_bytes
from app.pipeline import analyze_account
from app.providers.mock import MockProvider


def test_pipeline_success():
    row = analyze_account("yoga_bijo", MockProvider())
    assert row.status == "成功"
    assert row.followers_count == 48200
    assert row.metrics.engagement_rate is not None
    assert row.metrics.feed_avg_likes is not None
    assert row.metrics.reel_avg_views is not None
    assert len(row.metrics.top_posts) > 0
    assert row.profile_text


def test_pipeline_error_isolated():
    # one bad account should not stop the rest
    results = [analyze_account(u, MockProvider())
               for u in ["yoga_bijo", "nonexistent_user", "muscle_taro"]]
    assert results[0].status == "成功"
    assert results[1].status == "エラー"
    assert results[1].error_reason
    assert results[2].status == "成功"


def test_error_reasons_are_japanese_only():
    """Meta returns English errors (e.g. 'Invalid user id') — the Excel must never
    show them. Every mapped reason must contain Japanese and no English words."""
    import re

    import httpx

    from app.providers.graph_api import GraphAPIProvider

    def fake(code, message, **extra):
        payload = {"error": {"code": code, "message": message, **extra}}
        return httpx.Response(400, json=payload)

    cases = [
        fake(110, "Invalid user id",
             error_user_msg="ユーザーネームがxのユーザーが見つかりません"),
        fake(803, "Some node does not exist"),
        fake(190, "Error validating access token: Session has expired"),
        fake(4, "Application request limit reached"),
        fake(999, "Totally unknown english error"),  # unmapped -> must still be JP
    ]
    for resp in cases:
        reason = GraphAPIProvider._map_error(resp)
        assert re.search(r"[ぁ-んァ-ン一-龥]", reason), f"not Japanese: {reason}"
        assert not re.search(r"[A-Za-z]{4,}", reason), f"English leaked: {reason}"


def test_excel_structure():
    rows = [analyze_account(u, MockProvider())
            for u in ["yoga_bijo", "private_account"]]
    data = build_excel_bytes(rows)
    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    header_row = [c.value for c in ws[2]]
    assert header_row == HEADERS
    # expanded エーストリーム column set (identity+demographics+metrics+top5+profile+status)
    assert "直近12投稿のリール平均再生数" in HEADERS
    assert "人気投稿1（直近6ヶ月）" in HEADERS
    assert "プロフィール" in HEADERS
    # 1 note + 1 header + 2 data rows
    assert ws.max_row == 4
